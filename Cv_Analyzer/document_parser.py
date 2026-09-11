import os
import fitz
from docx import Document
import os
# Fix PaddleOCR NotImplementedError
os.environ["FLAGS_use_mkldnn"] = "0"
os.environ["PADDLE_PDX_ENABLE_MKLDNN_BYDEFAULT"] = "0"

from paddleocr import PaddleOCR


# Initialize OCR once
ocr = PaddleOCR(
    lang="en"
)


def extract_from_pdf(file_path):
    """
    Extract text from a normal PDF.
    """

    doc = fitz.open(file_path)

    text = ""

    for page in doc:
        page_text = page.get_text()

        if page_text:
            text += page_text + "\n"

    doc.close()

    return text


def extract_from_docx(file_path):
    """
    Extract text from DOCX.
    """

    doc = Document(file_path)

    text = ""

    # Normal paragraphs
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"

    # Tables
    for table in doc.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:
                row_text.append(cell.text.strip())

            text += " | ".join(row_text) + "\n"

    return text


def extract_from_image(file_path):
    """
    Extract text from an image using OCR.
    """

    result = ocr.predict(file_path)

    text_lines = []

    for res in result:

        if hasattr(res, "json"):
            data = res.json

            if callable(data):
                data = data()

            # PaddleOCR result structure can vary by version
            if isinstance(data, str):
                import json
                data = json.loads(data)

            if isinstance(data, dict):

                ocr_data = data.get("res", data)

                texts = ocr_data.get("rec_texts", [])
                scores = ocr_data.get("rec_scores", [])

                for i, line in enumerate(texts):

                    line = str(line).strip()

                    if not line:
                        continue

                    # Keep only reasonably confident OCR results
                    if scores:
                        try:
                            score = float(scores[i])

                            if score < 0.30:
                                continue

                        except (ValueError, IndexError):
                            pass

                    text_lines.append(line)

    return "\n".join(text_lines)

    return text


def extract_from_file(file_path):
    """
    Automatically choose the correct extraction method
    based on the file extension.
    """

    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".pdf":

        text = extract_from_pdf(file_path)

        # If PDF contains little/no text,
        # treat it as a scanned PDF and use OCR.
        if len(text.strip()) < 50:

            print("PDF appears to be scanned.")
            print("Running OCR...")

            doc = fitz.open(file_path)

            text = ""

            for page_number, page in enumerate(doc):

                pix = page.get_pixmap(
                    matrix=fitz.Matrix(2, 2)
                )

                image_path = f"_temp_page_{page_number}.png"

                pix.save(image_path)

                text += extract_from_image(image_path)
                text += "\n"

                os.remove(image_path)

            doc.close()

        return text

    elif extension == ".docx":

        return extract_from_docx(file_path)

    elif extension in [".jpg", ".jpeg", ".png"]:

        return extract_from_image(file_path)

    else:

        raise ValueError(
            "Unsupported file type. "
            "Use PDF, DOCX, JPG, JPEG, or PNG."
        )